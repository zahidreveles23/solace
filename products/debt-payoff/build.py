"""Builds Debt Free Plan: snowball vs avalanche debt payoff planner (Excel + Google Sheets)."""
import sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

DEMO = "--demo" in sys.argv
OUT = sys.argv[1]

NAVY, TEAL, CORAL, LIGHT, INPUT, GREY = "1F3A5F", "2A9D8F", "E76F51", "F3F6F9", "FFF4D6", "6B7785"
def font(**kw): return Font(name="Arial", **kw)
TITLE, SUB = font(size=20, bold=True, color=NAVY), font(size=10, italic=True, color=GREY)
HEAD, BODY, BOLD = font(size=10, bold=True, color="FFFFFF"), font(size=10), font(size=10, bold=True)
BLUE = font(size=10, color="0000FF")
HFILL, TFILL = PatternFill("solid", fgColor=NAVY), PatternFill("solid", fgColor=TEAL)
LFILL, IFILL = PatternFill("solid", fgColor=LIGHT), PatternFill("solid", fgColor=INPUT)
thin = Side(style="thin", color="D5DBE3"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD, USD0, PCT = '$#,##0.00;($#,##0.00);"-"', '$#,##0;($#,##0);"-"', '0.00%'
MONTHF = "mmm yyyy"
ND = 10              # max debts
NM = 360             # months simulated (30 years)
D0 = 6               # first debt row on My Debts
DR = f"${D0}:${D0+ND-1}"

wb = Workbook()
def setup(ws, title, subtitle, widths):
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font, ws["A2"].font = TITLE, SUB
    for col, w in widths.items(): ws.column_dimensions[col].width = w
def head(ws, row, col, labels, fill=HFILL):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill, c.border = HEAD, fill, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
def cell(ws, ref, value, fmt=None, f=BODY, fill=None, border=True):
    c = ws[ref]; c.value = value; c.font = f
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if border: c.border = BOX
    return c

# ------------------------------------------------------------ Start Here
st = wb.active; st.title = "Start Here"
setup(st, "Debt Free Plan", "See your debt-free date and exactly what to pay each month", {"A": 3, "B": 40, "C": 18, "D": 62})
st["B4"] = "YOUR SETTINGS"; st["B4"].font = font(size=12, bold=True, color=TEAL)
settings = [("Total you can pay toward debt each month", 1200 if not DEMO else 850, USD0, "Include all minimum payments plus any extra you can afford."),
            ("Plan start month", date(2026, 10, 1), MONTHF, "The first month you'll follow this plan."),
            ("Payoff method", "Avalanche", None, "Avalanche = highest interest first (saves the most). Snowball = smallest balance first (fastest wins).")]
for i, (lab, val, fmt, note) in enumerate(settings):
    r = 5 + i
    cell(st, f"B{r}", lab, border=False)
    cell(st, f"C{r}", val, fmt, BLUE, IFILL)
    st[f"D{r}"] = note; st[f"D{r}"].font = SUB
BUDGET, START, METHOD = "'Start Here'!$C$5", "'Start Here'!$C$6", "'Start Here'!$C$7"
v = DataValidation(type="list", formula1='"Avalanche,Snowball"', allow_blank=False); st.add_data_validation(v); v.add("C7")
steps = ["1.  List your debts on the My Debts tab: name, balance, interest rate (APR) and minimum payment.",
         "2.  Enter your monthly debt budget above (yellow cells = type here).",
         "3.  Open My Plan to see your debt-free date, total interest, and what to pay this month.",
         "4.  Pick Avalanche or Snowball above. My Plan compares both so you can see the difference.",
         "5.  Each month, pay the amounts shown. When one debt is paid off, its payment rolls to the next.",
         "6.  Update balances every few months so the plan stays accurate.",
         "Tip: in Google Sheets, open the file from Google Drive or use File > Import."]
st["B10"] = "HOW TO USE"; st["B10"].font = font(size=12, bold=True, color=TEAL)
for i, s in enumerate(steps): st.cell(row=11 + i, column=2, value=s).font = BODY
st["B19"] = "Estimates assume fixed interest rates and on-time payments, with no new charges. Not financial advice."
st["B19"].font = SUB

# ------------------------------------------------------------ My Debts
dt = wb.create_sheet("My Debts")
setup(dt, "My Debts", "Enter up to 10 debts. Yellow cells are yours; grey cells calculate automatically.",
      {"A": 3, "B": 28, "C": 15, "D": 12, "E": 15, "F": 13, "G": 13, "H": 16, "I": 16})
head(dt, 5, 2, ["Debt name", "Balance ($)", "APR", "Minimum payment ($)", "Snowball order", "Avalanche order", "Paid off (your plan)", "Interest paid (your plan)"])
demo = [("Store card", 850, 0.2699, 35), ("Visa", 4200, 0.2249, 110), ("Car loan", 9800, 0.069, 285), ("Personal loan", 3100, 0.1199, 120), ("Medical bill", 600, 0.0, 50)]
example = demo[:3]
bal, apr, mp = (f"'My Debts'!$C${D0}:$C${D0+ND-1}", f"'My Debts'!$D${D0}:$D${D0+ND-1}", f"'My Debts'!$E${D0}:$E${D0+ND-1}")
for i in range(ND):
    r = D0 + i
    for col, fmt in (("B", None), ("C", USD), ("D", PCT), ("E", USD)):
        cell(dt, f"{col}{r}", None, fmt, BLUE, IFILL)
    rows = demo if DEMO else example
    if i < len(rows):
        n, b, a, m = rows[i]
        dt[f"B{r}"], dt[f"C{r}"], dt[f"D{r}"], dt[f"E{r}"] = (n if DEMO else f"{n} (example)"), b, a, m
        if not DEMO:
            for col in "BCDE": dt[f"{col}{r}"].font = font(size=10, italic=True, color=GREY)
    cell(dt, f"F{r}", f'=IF(C{r}="",99,SUMPRODUCT(--($C${D0}:$C${D0+ND-1}<>""),--($C${D0}:$C${D0+ND-1}<C{r}))+COUNTIF($C${D0}:C{r},C{r}))', "0", fill=LFILL)
    cell(dt, f"G{r}", f'=IF(C{r}="",99,SUMPRODUCT(--($C${D0}:$C${D0+ND-1}<>""),--($D${D0}:$D${D0+ND-1}>D{r}))+COUNTIF($D${D0}:D{r},D{r}))', "0", fill=LFILL)
    for c in "FG":
        dt.conditional_formatting.add(f"{c}{r}", CellIsRule(operator="equal", formula=["99"], font=font(size=10, color="FFFFFF")))
T = D0 + ND
cell(dt, f"B{T}", "Total", f=BOLD, fill=LFILL)
for col, fmt in (("C", USD), ("E", USD)):
    cell(dt, f"{col}{T}", f"=SUM({col}{D0}:{col}{T-1})", fmt, BOLD, LFILL)
cell(dt, f"D{T}", f'=IF(C{T}>0,SUMPRODUCT(C{D0}:C{T-1},D{D0}:D{T-1})/C{T},0)', PCT, BOLD, LFILL)
dt[f"D{T}"].comment = Comment("Balance-weighted average APR.", "Debt Free Plan")
for col in "FGHI": cell(dt, f"{col}{T}", None, fill=LFILL)
dt[f"B{T+2}"] = "Delete the example rows and enter your own debts. Rows left blank are ignored."
dt[f"B{T+2}"].font = SUB
dt["D5"].comment = Comment("Annual interest rate from your statement, e.g. 22.49%. Enter 0% for interest-free debts.", "Debt Free Plan")

# ------------------------------------------------------------ Schedules
def schedule(name, prio_col):
    ws = wb.create_sheet(name)
    setup(ws, f"{name.split()[0]} Schedule", "Month-by-month simulation. No need to edit this tab.", {"A": 11})
    # blocks: need B.., min, end
    NB, MB, EB = 2, 2 + ND, 2 + 2 * ND
    POOL, TOT, INT = 2 + 3 * ND, 3 + 3 * ND, 4 + 3 * ND
    ws.cell(row=4, column=1, value="Priority").font = BOLD
    for i in range(ND):
        dr = D0 + i
        ws.cell(row=4, column=NB + i, value=f"='My Debts'!${prio_col}${dr}").font = BODY
        for blk, lab in ((NB, "Owed"), (MB, "Minimum"), (EB, "End balance")):
            c = ws.cell(row=5, column=blk + i, value=f'=IF(\'My Debts\'!$B${dr}="","Debt {i+1}",\'My Debts\'!$B${dr})&" · {lab}"')
            c.font, c.fill = HEAD, HFILL; c.alignment = Alignment(wrap_text=True, horizontal="center")
        ws.cell(row=6, column=EB + i, value=f"=N('My Debts'!$C${dr})").number_format = USD
        for blk in (NB, MB, EB): ws.column_dimensions[L(blk + i)].width = 12
    for col, lab in ((1, "Month"), (POOL, "Extra available"), (TOT, "Total balance"), (INT, "Interest this month")):
        c = ws.cell(row=5, column=col, value=lab); c.font, c.fill = HEAD, HFILL; c.alignment = Alignment(wrap_text=True, horizontal="center")
        ws.column_dimensions[L(col)].width = 13
    ws.row_dimensions[5].height = 42
    ws.cell(row=6, column=1, value="Start").font = BOLD
    ws.cell(row=6, column=TOT, value=f"=SUM({L(EB)}6:{L(EB+ND-1)}6)").number_format = USD
    nr = lambda r: f"${L(NB)}{r}:${L(NB+ND-1)}{r}"
    mr = lambda r: f"${L(MB)}{r}:${L(MB+ND-1)}{r}"
    prio = f"${L(NB)}$4:${L(NB+ND-1)}$4"
    for m in range(1, NM + 1):
        r = 6 + m
        ws.cell(row=r, column=1, value=f"=DATE(YEAR({START}),MONTH({START})+{m-1},1)").number_format = MONTHF
        for i in range(ND):
            dr = D0 + i
            n, mn, e = L(NB + i), L(MB + i), L(EB + i)
            ws[f"{n}{r}"] = f"=ROUND({e}{r-1}*(1+N('My Debts'!$D${dr})/12),2)"
            ws[f"{mn}{r}"] = f"=MIN(N('My Debts'!$E${dr}),{n}{r})"
            ws[f"{e}{r}"] = (f"=ROUND({n}{r}-{mn}{r}-MIN({n}{r}-{mn}{r},MAX(0,${L(POOL)}{r}"
                             f"-SUMPRODUCT(--({prio}<{n}$4),{nr(r)}-{mr(r)}))),2)")
            for c in (n, mn, e): ws[f"{c}{r}"].number_format = USD
        ws.cell(row=r, column=POOL, value=f"=MAX(0,{BUDGET}-SUM({mr(r)}))").number_format = USD
        ws.cell(row=r, column=TOT, value=f"=SUM({L(EB)}{r}:{L(EB+ND-1)}{r})").number_format = USD
        ws.cell(row=r, column=INT, value=f"=SUM({L(NB)}{r}:{L(NB+ND-1)}{r})-{L(TOT)}{r-1}").number_format = USD
    ws.freeze_panes = "B6"
    last = 6 + NM
    return {"ws": ws, "EB": EB, "NB": NB, "TOT": L(TOT), "INT": L(INT), "last": last}

SN = schedule("Snowball Schedule", "F")
AV = schedule("Avalanche Schedule", "G")

def sref(s, col_letter, r0=7): return f"'{s['ws'].title}'!${col_letter}${r0}:${col_letter}${s['last']}"
def months_to_free(s): return f"COUNTIF({sref(s, s['TOT'])},\">0.005\")+1"
def interest(s): return f"SUM({sref(s, s['INT'])})"

# My Debts: payoff + interest per debt under chosen plan
for i in range(ND):
    r = D0 + i
    per = []
    for s in (SN, AV):
        e = L(s["EB"] + i)
        per.append(f"COUNTIF({sref(s, e)},\">0.005\")+1")
    months = f'IF({METHOD}="Snowball",{per[0]},{per[1]})'
    cell(dt, f"H{r}", f'=IF(C{r}="","",IF({months}>{NM},"30+ years",DATE(YEAR({START}),MONTH({START})+{months}-1,1)))', MONTHF, fill=LFILL)
    ints = []
    for s in (SN, AV):
        n, e = L(s["NB"] + i), L(s["EB"] + i)
        t = s["ws"].title
        ints.append(f"SUM('{t}'!${n}$7:${n}${s['last']})-SUM('{t}'!${e}$6:${e}${s['last']-1})")
    cell(dt, f"I{r}", f'=IF(C{r}="","",IF({METHOD}="Snowball",{ints[0]},{ints[1]}))', USD, fill=LFILL)

# ------------------------------------------------------------ My Plan
pl = wb.create_sheet("My Plan", 1)
setup(pl, "My Plan", "", {"A": 3, "B": 26, "C": 16, "D": 16, "E": 16, "F": 3, "G": 24, "H": 16, "I": 16, "J": 16})
pl["B2"] = f'="Your "&LOWER({METHOD})&" plan · starting "&TEXT({START},"mmmm yyyy")'
pl["B2"].font = font(size=11, bold=True, color=GREY)
pl["A2"] = None
# warning
pl["B3"] = f'=IF({BUDGET}<\'My Debts\'!$E${T},"⚠ Your monthly budget is less than your minimum payments ("&TEXT(\'My Debts\'!$E${T},"$#,##0")&"). This plan assumes you still pay every minimum. Raise your budget on Start Here.","")'
pl["B3"].font = font(size=11, bold=True, color="C0392B")

chosen_m = f'IF({METHOD}="Snowball",{months_to_free(SN)},{months_to_free(AV)})'
chosen_i = f'IF({METHOD}="Snowball",{interest(SN)},{interest(AV)})'
kpis = [("B", "Debt-free date", f'=IF({chosen_m}>{NM},"30+ years",DATE(YEAR({START}),MONTH({START})+{chosen_m}-1,1))', MONTHF),
        ("D", "Months to go", f'=IF({chosen_m}>{NM},"30+",{chosen_m})', "0"),
        ("G", "Total interest", f"={chosen_i}", USD0),
        ("I", "Total you'll pay", f"='My Debts'!$C${T}+{chosen_i}", USD0)]
for col, lab, f, fmt in kpis:
    c2 = chr(ord(col) + 1)
    pl.merge_cells(f"{col}5:{c2}5"); pl.merge_cells(f"{col}6:{c2}7")
    pl[f"{col}5"] = lab.upper(); pl[f"{col}5"].font = font(size=9, bold=True, color="FFFFFF"); pl[f"{col}5"].fill = TFILL
    pl[f"{col}5"].alignment = Alignment(horizontal="center")
    pl[f"{col}6"] = f; pl[f"{col}6"].number_format = fmt
    pl[f"{col}6"].font = font(size=20, bold=True, color=NAVY); pl[f"{col}6"].alignment = Alignment(horizontal="center", vertical="center")
    for rr in (5, 6, 7):
        for cc in (col, c2):
            pl[f"{cc}{rr}"].border = BOX
            if rr > 5: pl[f"{cc}{rr}"].fill = LFILL

# Pay this month
pl["B10"] = "WHAT TO PAY THIS MONTH"; pl["B10"].font = font(size=12, bold=True, color=TEAL)
pl["D10"] = "Month #"; pl["D10"].font = BOLD; pl["D10"].alignment = Alignment(horizontal="right")
cell(pl, "E10", 1, "0", BLUE, IFILL)
pl["E10"].comment = Comment("Change to see any future month (1 = first month of the plan).", "Debt Free Plan")
head(pl, 11, 2, ["Debt", "Pay this month", "Balance after", "Order"])
for i in range(ND):
    r = 12 + i
    dr = D0 + i
    pays, ends = [], []
    for s in (SN, AV):
        t = s["ws"].title; n, e = L(s["NB"] + i), L(s["EB"] + i)
        pays.append(f"INDEX('{t}'!${n}$7:${n}${s['last']},$E$10)-INDEX('{t}'!${e}$7:${e}${s['last']},$E$10)")
        ends.append(f"INDEX('{t}'!${e}$7:${e}${s['last']},$E$10)")
    cell(pl, f"B{r}", f"=IF('My Debts'!$B${dr}=\"\",\"\",'My Debts'!$B${dr})")
    cell(pl, f"C{r}", f'=IF(B{r}="","",IF({METHOD}="Snowball",{pays[0]},{pays[1]}))', USD)
    cell(pl, f"D{r}", f'=IF(B{r}="","",IF({METHOD}="Snowball",{ends[0]},{ends[1]}))', USD)
    cell(pl, f"E{r}", f"=IF(B{r}=\"\",\"\",IF({METHOD}=\"Snowball\",'My Debts'!$F${dr},'My Debts'!$G${dr}))", "0")
cell(pl, f"B{12+ND}", "Total", f=BOLD, fill=LFILL)
for col in "CD": cell(pl, f"{col}{12+ND}", f"=SUM({col}12:{col}{11+ND})", USD, BOLD, LFILL)
cell(pl, f"E{12+ND}", None, fill=LFILL)

# Comparison
pl["G10"] = "SNOWBALL vs AVALANCHE"; pl["G10"].font = font(size=12, bold=True, color=TEAL)
head(pl, 11, 7, ["", "Snowball", "Avalanche", "Difference"])
comp = [("Debt-free in (months)", months_to_free(SN), months_to_free(AV), "0"),
        ("Total interest", interest(SN), interest(AV), USD0)]
for k, (lab, a, b, fmt) in enumerate(comp):
    r = 12 + k
    cell(pl, f"G{r}", lab, f=BOLD)
    if k == 0:
        cell(pl, f"H{r}", f'=IF({a}>{NM},"30+ yrs",{a})', fmt)
        cell(pl, f"I{r}", f'=IF({b}>{NM},"30+ yrs",{b})', fmt)
    else:
        cell(pl, f"H{r}", f"={a}", fmt)
        cell(pl, f"I{r}", f"={b}", fmt)
    cell(pl, f"J{r}", f'=IFERROR(H{r}-I{r},"-")', fmt)
pl["G15"] = (f'=IF(\'My Debts\'!$C${T}=0,"",IF(J13>0.5,"Avalanche saves you "&TEXT(J13,"$#,##0")&" in interest.",'
             f'"Both methods cost about the same here, so Snowball\'s quick wins are a free bonus."))')
pl["G15"].font = font(size=11, bold=True, color=NAVY)
pl["G16"] = "Snowball pays off small debts first for quick wins and motivation."; pl["G16"].font = SUB
pl["G17"] = "Avalanche targets the highest interest first and usually costs less."; pl["G17"].font = SUB

# Payoff order
pl["G19"] = "PAYOFF ORDER (YOUR PLAN)"; pl["G19"].font = font(size=12, bold=True, color=TEAL)
head(pl, 20, 7, ["#", "Debt", "Paid off", "Interest paid"])
for k in range(1, ND + 1):
    r = 20 + k
    pcol = f'IF({METHOD}="Snowball",\'My Debts\'!$F${D0}:$F${D0+ND-1},\'My Debts\'!$G${D0}:$G${D0+ND-1})'
    idx = f"MATCH({k},{pcol},0)"
    cell(pl, f"G{r}", f'=IF(ISNA({idx}),"",{k})', "0")
    cell(pl, f"H{r}", f'=IF(G{r}="","",INDEX(\'My Debts\'!$B${D0}:$B${D0+ND-1},{idx}))')
    cell(pl, f"I{r}", f'=IF(G{r}="","",INDEX(\'My Debts\'!$H${D0}:$H${D0+ND-1},{idx}))', MONTHF)
    cell(pl, f"J{r}", f'=IF(G{r}="","",INDEX(\'My Debts\'!$I${D0}:$I${D0+ND-1},{idx}))', USD0)

# Chart: 41 points spread across the longer plan, so any payoff length fills the chart
CH = 34
pl[f"B{CH-1}"] = "BALANCE OVER TIME"; pl[f"B{CH-1}"].font = font(size=12, bold=True, color=TEAL)
cd = AV["ws"]; C0 = 6 + 3 * ND  # helper columns to the right of the avalanche schedule
for j, lab in enumerate(["Chart: month", "Snowball", "Avalanche"]):
    c = cd.cell(row=5, column=C0 + j, value=lab); c.font, c.fill = HEAD, TFILL
    cd.column_dimensions[L(C0 + j)].width = 12
step = f"MAX(1,CEILING(MAX({months_to_free(SN)},{months_to_free(AV)})/40,1))"
cd.cell(row=4, column=C0, value="Step").font = BOLD
cd.cell(row=4, column=C0 + 1, value=f"=MIN(9,{step})")
for k in range(41):
    r = 6 + k
    cd.cell(row=r, column=C0, value=f"={k}*${L(C0+1)}$4")
    cd.cell(row=r, column=C0 + 1, value=f"=INDEX({sref(SN, SN['TOT'], 6)},{L(C0)}{r}+1)").number_format = USD0
    cd.cell(row=r, column=C0 + 2, value=f"=INDEX({sref(AV, AV['TOT'], 6)},{L(C0)}{r}+1)").number_format = USD0
lc = LineChart(); lc.title = "Total debt balance"; lc.height = 9; lc.width = 24
lc.add_data(Reference(cd, min_col=C0 + 1, max_col=C0 + 2, min_row=5, max_row=46), titles_from_data=True)
lc.set_categories(Reference(cd, min_col=C0, min_row=6, max_row=46))
lc.y_axis.numFmt = "$#,##0"; lc.x_axis.title = "Months from start"; lc.y_axis.majorGridlines = None
lc.x_axis.tickLblSkip = 5; lc.x_axis.delete = False; lc.y_axis.delete = False
lc.series[0].graphicalProperties.line.solidFill = CORAL; lc.series[1].graphicalProperties.line.solidFill = TEAL
for s_ in lc.series: s_.graphicalProperties.line.width = 28000; s_.smooth = False
pl.add_chart(lc, f"B{CH+1}")

for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
for s in (SN, AV): s["ws"].sheet_properties.tabColor = "B0B8C4"
pl.page_setup.fitToHeight = 1
wb.active = 0
wb.save(OUT); print("saved", OUT, wb.sheetnames)
