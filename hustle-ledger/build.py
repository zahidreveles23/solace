"""Builds Hustle Ledger: side-hustle income, expense & tax tracker (Excel + Google Sheets)."""
import random, sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule, CellIsRule
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment

DEMO = "--demo" in sys.argv
OUT = sys.argv[1]

NAVY, TEAL, LIGHT, INPUT, GREY = "1F3A5F", "2A9D8F", "F3F6F9", "FFF4D6", "6B7785"
F = "Arial"
def font(**kw): return Font(name=F, **kw)
TITLE = font(size=20, bold=True, color=NAVY)
SUB = font(size=10, italic=True, color=GREY)
HEAD = font(size=10, bold=True, color="FFFFFF")
BODY = font(size=10)
BOLD = font(size=10, bold=True)
HFILL = PatternFill("solid", fgColor=NAVY)
TFILL = PatternFill("solid", fgColor=TEAL)
LFILL = PatternFill("solid", fgColor=LIGHT)
IFILL = PatternFill("solid", fgColor=INPUT)
thin = Side(style="thin", color="D5DBE3")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATEF = "mmm d, yyyy"
N = 500                     # data rows per log
R0, R1 = 5, 5 + N - 1       # first/last data row

wb = Workbook()

def setup(ws, title, subtitle, widths):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title; ws["A1"].font = TITLE
    ws["A2"] = subtitle; ws["A2"].font = SUB
    for col, w in widths.items(): ws.column_dimensions[col].width = w

def header(ws, row, labels, fill=HFILL):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.border = HEAD, fill, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30

# ---------------------------------------------------------------- Start Here
st = wb.active; st.title = "Start Here"
setup(st, "Hustle Ledger", "Side-hustle & freelance income, expense and tax tracker", {"A": 3, "B": 44, "C": 18, "D": 60})
st["B4"] = "YOUR SETTINGS"; st["B4"].font = font(size=12, bold=True, color=TEAL)
settings = [
    ("Business name", "My Side Hustle", None, "Shown on the Dashboard."),
    ("Tax year", 2026, "0", "Only entries dated in this year are counted."),
    ("Monthly income goal", 3000, USD0, "Used for goal progress on the Dashboard."),
    ("Federal income tax rate (estimate)", 0.12, PCT, "Your expected federal bracket, e.g. 10%, 12%, 22%."),
    ("State income tax rate (estimate)", 0.05, PCT, "Use 0% if your state has no income tax."),
    ("Self-employment tax rate", 0.153, PCT, "US SE tax: 12.4% Social Security + 2.9% Medicare."),
    ("Share of profit subject to SE tax", 0.9235, PCT, "IRS Schedule SE uses 92.35% of net profit."),
    ("Mileage rate ($ per mile)", 0.70, '$0.000', "IRS standard rate for 2025 was $0.70. Update to the current year's rate at irs.gov."),
]
for i, (lab, val, fmt, note) in enumerate(settings):
    r = 5 + i
    st.cell(row=r, column=2, value=lab).font = BODY
    c = st.cell(row=r, column=3, value=val); c.font = font(size=10, color="0000FF"); c.fill = IFILL; c.border = BOX
    if fmt: c.number_format = fmt
    st.cell(row=r, column=4, value=note).font = SUB
S = {k: f"'Start Here'!$C${5+i}" for i, k in enumerate(["name", "year", "goal", "fed", "state", "se", "seshare", "mile"])}

r = 15
st.cell(row=r, column=2, value="HOW TO USE").font = font(size=12, bold=True, color=TEAL)
steps = [
    "1.  Fill in the yellow cells above. Yellow = type here. Everything else calculates itself.",
    "2.  Log every payment you receive on the Income tab (date, source, amount).",
    "3.  Log every business purchase on the Expenses tab. Pick a category from the dropdown.",
    "4.  Log business driving on the Mileage tab. Deductions are calculated for you.",
    "5.  Check the Dashboard for monthly profit, goal progress and where your money goes.",
    "6.  Check Tax Estimate for how much to set aside and your quarterly payment amounts.",
    "7.  Delete the example row on each log tab (row 5) when you add your own data.",
    "Tip: in Google Sheets, use File > Import > Upload, or open the file from Google Drive.",
    "Tip: edit category names on the Lists tab. The dropdowns and Dashboard update automatically.",
]
for i, s in enumerate(steps):
    st.cell(row=r + 1 + i, column=2, value=s).font = BODY
r = 27
st.cell(row=r, column=2, value="IMPORTANT").font = font(size=12, bold=True, color="C0392B")
st.cell(row=r + 1, column=2, value=(
    "Tax figures are simplified estimates for planning only, not tax advice. They do not include the "
    "standard deduction, QBI deduction, credits, other income or the Social Security wage cap. "
    "Confirm amounts with a tax professional or IRS Form 1040-ES.")).font = BODY
st.merge_cells(start_row=r + 1, start_column=2, end_row=r + 3, end_column=4)
st.cell(row=r + 1, column=2).alignment = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------- Lists
ls = wb.create_sheet("Lists")
setup(ls, "Lists", "Edit these to customize dropdowns. Keep deductible % between 0% and 100%.", {"A": 26, "B": 3, "C": 26, "D": 14, "E": 3, "F": 22})
header(ls, 4, ["Income categories"]); ls["B4"].fill = PatternFill()
for col, lab in (("C", "Expense categories"), ("D", "Deductible %"), ("F", "Payment methods")):
    c = ls[f"{col}4"]; c.value = lab; c.font, c.fill, c.border = HEAD, HFILL, BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
inc_cats = ["Client work", "Product sales", "Services", "Tips", "Affiliate / ads", "Rental", "Other income"]
exp_cats = [("Advertising", 1), ("Software & subscriptions", 1), ("Supplies & materials", 1), ("Equipment", 1),
            ("Phone & internet", 1), ("Home office", 1), ("Platform & payment fees", 1), ("Shipping", 1),
            ("Education & training", 1), ("Travel", 1), ("Meals (business)", 0.5), ("Professional services", 1),
            ("Insurance", 1), ("Other", 1)]
pays = ["Bank transfer", "PayPal", "Stripe", "Cash", "Card", "Venmo", "Zelle", "Check", "Platform payout"]
LN = 25  # list capacity
for i in range(LN):
    for col in ("A", "C", "D", "F"):
        c = ls[f"{col}{5+i}"]; c.fill, c.border, c.font = IFILL, BOX, font(size=10, color="0000FF")
    ls[f"D{5+i}"].number_format = "0%"
    if i < len(inc_cats): ls[f"A{5+i}"] = inc_cats[i]
    if i < len(exp_cats): ls[f"C{5+i}"], ls[f"D{5+i}"] = exp_cats[i]
    if i < len(pays): ls[f"F{5+i}"] = pays[i]
ls["D15"].comment = Comment("Business meals are generally 50% deductible (IRS Publication 463).", "Hustle Ledger")
INC_L = f"Lists!$A$5:$A${4+LN}"; EXP_L = f"Lists!$C$5:$C${4+LN}"; PCT_L = f"Lists!$D$5:$D${4+LN}"; PAY_L = f"Lists!$F$5:$F${4+LN}"

def dv(ws, formula, rng):
    v = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(v); v.add(rng)

def input_rows(ws, cols, calc_cols, fmts):
    for r in range(R0, R1 + 1):
        for ci, col in enumerate(cols):
            c = ws[f"{col}{r}"]; c.border = BOX; c.font = BODY
            if col in fmts: c.number_format = fmts[col]
            if col in calc_cols: c.fill = LFILL
    ws.freeze_panes = f"A{R0}"

# ---------------------------------------------------------------- Income
inc = wb.create_sheet("Income")
setup(inc, "Income", "Log every payment you receive. Type in the white cells; totals update everywhere automatically.", {"A": 14, "B": 26, "C": 34, "D": 20, "E": 14, "F": 18})
header(inc, 4, ["Date", "Client / source", "Description", "Category", "Amount ($)", "Paid via"])
input_rows(inc, "ABCDEF", "", {"A": DATEF, "E": USD})
dv(inc, "=" + INC_L, f"D{R0}:D{R1}"); dv(inc, "=" + PAY_L, f"F{R0}:F{R1}")
inc["A3"] = "Total income logged (all years):"; inc["A3"].font = BOLD
inc["E3"] = f"=SUM(E{R0}:E{R1})"; inc["E3"].font = BOLD; inc["E3"].number_format = USD
inc.auto_filter.ref = f"A4:F{R1}"

# ---------------------------------------------------------------- Expenses
exp = wb.create_sheet("Expenses")
setup(exp, "Expenses", "Log business purchases. Grey cells calculate the deductible amount for you.", {"A": 14, "B": 24, "C": 32, "D": 26, "E": 14, "F": 14, "G": 16})
header(exp, 4, ["Date", "Vendor", "Description", "Category", "Amount ($)", "Business use %", "Deductible ($)"])
exp["F4"].comment = Comment("Leave blank for 100%. For mixed-use items (e.g. a phone used 60% for business) enter 60%.", "Hustle Ledger")
input_rows(exp, "ABCDEFG", "G", {"A": DATEF, "E": USD, "F": "0%", "G": USD})
for r in range(R0, R1 + 1):
    exp[f"G{r}"] = (f'=IF(E{r}="","",E{r}*IFERROR(INDEX({PCT_L},MATCH(D{r},{EXP_L},0)),1)'
                    f'*IF(F{r}="",1,F{r}))')
dv(exp, "=" + EXP_L, f"D{R0}:D{R1}")
exp["A3"] = "Totals (all years):"; exp["A3"].font = BOLD
for col in "EG":
    exp[f"{col}3"] = f"=SUM({col}{R0}:{col}{R1})"; exp[f"{col}3"].font = BOLD; exp[f"{col}3"].number_format = USD
exp.auto_filter.ref = f"A4:G{R1}"

# ---------------------------------------------------------------- Mileage
mi = wb.create_sheet("Mileage")
setup(mi, "Mileage", "Log business trips (not your commute). Grey cells calculate the deduction for you.", {"A": 14, "B": 22, "C": 22, "D": 30, "E": 12, "F": 16})
header(mi, 4, ["Date", "From", "To", "Purpose", "Miles", "Deduction ($)"])
input_rows(mi, "ABCDEF", "F", {"A": DATEF, "E": "#,##0.0", "F": USD})
for r in range(R0, R1 + 1):
    mi[f"F{r}"] = f'=IF(E{r}="","",E{r}*{S["mile"]})'
mi["A3"] = "Totals (all years):"; mi["A3"].font = BOLD
for col, fmt in (("E", "#,##0.0"), ("F", USD)):
    mi[f"{col}3"] = f"=SUM({col}{R0}:{col}{R1})"; mi[f"{col}3"].font = BOLD; mi[f"{col}3"].number_format = fmt

# ---------------------------------------------------------------- Dashboard
db = wb.create_sheet("Dashboard", 1)
setup(db, "Dashboard", "", {"A": 3, "B": 16, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 3, "J": 26, "K": 15, "L": 12})
db["B2"] = f'={S["name"]}&" · "&{S["year"]}&" overview"'; db["B2"].font = font(size=11, bold=True, color=GREY)
db["A2"] = None

def rng(ws, col): return f"{ws}!${col}${R0}:${col}${R1}"
def in_month(ws, datecol, valcol, m):
    y = S["year"]
    return (f'SUMIFS({rng(ws, valcol)},{rng(ws, datecol)},">="&DATE({y},{m},1),'
            f'{rng(ws, datecol)},"<"&DATE({y},{m}+1,1))')

# Monthly table
T0 = 12
db.cell(row=T0 - 1, column=2, value="MONTH BY MONTH").font = font(size=12, bold=True, color=TEAL)
cols = ["Month", "Income", "Expenses", "Deductions", "Taxable profit", "Tax to set aside", "Take-home", "Goal progress"]
for i, lab in enumerate(cols):
    c = db.cell(row=T0, column=2 + i, value=lab); c.font, c.fill, c.border = HEAD, HFILL, BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
db.row_dimensions[T0].height = 30
months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
EFF = "'Tax Estimate'!$C$18"  # effective combined tax rate on profit
for m in range(1, 13):
    r = T0 + m
    db[f"B{r}"] = months[m - 1]
    db[f"C{r}"] = "=" + in_month("Income", "A", "E", m)
    db[f"D{r}"] = "=" + in_month("Expenses", "A", "E", m)
    db[f"E{r}"] = "=" + in_month("Expenses", "A", "G", m) + "+" + in_month("Mileage", "A", "F", m)
    db[f"F{r}"] = f"=C{r}-E{r}"
    db[f"G{r}"] = f"=MAX(0,F{r})*{EFF}"
    db[f"H{r}"] = f"=C{r}-D{r}-G{r}"
    db[f"I{r}"] = f'=IF({S["goal"]}>0,C{r}/{S["goal"]},0)'
TT = T0 + 13
db[f"B{TT}"] = "Total"
for col in "CDEFGH":
    db[f"{col}{TT}"] = f"=SUM({col}{T0+1}:{col}{T0+12})"
db[f"I{TT}"] = f'=IF({S["goal"]}>0,C{TT}/({S["goal"]}*12),0)'
for r in range(T0 + 1, TT + 1):
    for ci in range(2, 10):
        c = db.cell(row=r, column=ci); c.border = BOX
        c.font = BOLD if r == TT else BODY
        c.number_format = PCT if ci == 9 else USD0
        if r == TT: c.fill = LFILL
db.column_dimensions["I"].width = 14
db.conditional_formatting.add(f"I{T0+1}:I{T0+12}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=TEAL))
db.conditional_formatting.add(f"F{T0+1}:F{T0+12}", CellIsRule(operator="lessThan", formula=["0"], font=font(size=10, color="C0392B")))

# KPI tiles (row 4-8)
kpis = [("B", "Income", f"=C{TT}", USD0), ("D", "Expenses", f"=D{TT}", USD0),
        ("F", "Tax to set aside", f"=G{TT}", USD0), ("H", "Take-home", f"=H{TT}", USD0)]
for col, lab, f, fmt in kpis:
    c2 = chr(ord(col) + 1)
    db.merge_cells(f"{col}4:{c2}4"); db.merge_cells(f"{col}5:{c2}6")
    db[f"{col}4"] = lab.upper(); db[f"{col}4"].font = font(size=9, bold=True, color="FFFFFF")
    db[f"{col}4"].fill = TFILL; db[f"{col}4"].alignment = Alignment(horizontal="center")
    db[f"{col}5"] = f; db[f"{col}5"].font = font(size=20, bold=True, color=NAVY); db[f"{col}5"].number_format = fmt
    db[f"{col}5"].alignment = Alignment(horizontal="center", vertical="center")
    for r in (4, 5, 6):
        for cc in (col, c2):
            db[f"{cc}{r}"].border = BOX
            if r > 4: db[f"{cc}{r}"].fill = LFILL
db["B8"] = "Yearly goal progress"; db["B8"].font = BOLD
db["D8"] = f"=I{TT}"; db["D8"].number_format = PCT; db["D8"].font = font(size=12, bold=True, color=TEAL)
db["E8"] = f'="of "&TEXT({S["goal"]}*12,"$#,##0")&" goal"'; db["E8"].font = SUB
db["G8"] = "Profit margin"; db["G8"].font = BOLD
db["H8"] = f"=IF(C{TT}>0,(C{TT}-D{TT})/C{TT},0)"; db["H8"].number_format = PCT; db["H8"].font = font(size=12, bold=True, color=TEAL)

# Expenses by category
db.column_dimensions["J"].width = 3
db.column_dimensions["K"].width = 26; db.column_dimensions["L"].width = 14; db.column_dimensions["M"].width = 10
db["K11"] = "WHERE THE MONEY GOES"; db["K11"].font = font(size=12, bold=True, color=TEAL)
for col, lab in (("K", "Category"), ("L", "Spent"), ("M", "Share")):
    c = db[f"{col}12"]; c.value = lab; c.font, c.fill, c.border = HEAD, HFILL, BOX
    c.alignment = Alignment(horizontal="center", vertical="center")
y = S["year"]
for i in range(LN):
    r = 13 + i
    db[f"K{r}"] = f'=IF(Lists!$C${5+i}="","",Lists!$C${5+i})'
    db[f"L{r}"] = (f'=IF(K{r}="","",SUMIFS({rng("Expenses","E")},{rng("Expenses","D")},K{r},'
                   f'{rng("Expenses","A")},">="&DATE({y},1,1),{rng("Expenses","A")},"<"&DATE({y}+1,1,1)))')
    db[f"M{r}"] = f'=IF(OR(K{r}="",$D${TT}=0),"",L{r}/$D${TT})'
    for col in "KLM":
        db[f"{col}{r}"].border = BOX; db[f"{col}{r}"].font = BODY
    db[f"L{r}"].number_format = USD0; db[f"M{r}"].number_format = PCT
db.conditional_formatting.add(f"M13:M{12+LN}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="E9C46A"))

# Charts
bar = BarChart(); bar.type = "col"; bar.grouping = "clustered"
bar.title = "Income vs expenses by month"; bar.y_axis.title = None; bar.height = 8; bar.width = 22
bar.add_data(Reference(db, min_col=3, max_col=4, min_row=T0, max_row=T0 + 12), titles_from_data=True)
bar.set_categories(Reference(db, min_col=2, min_row=T0 + 1, max_row=T0 + 12))
bar.y_axis.numFmt = "$#,##0"; bar.y_axis.majorGridlines = None
bar.series[0].graphicalProperties.solidFill = TEAL; bar.series[1].graphicalProperties.solidFill = "E76F51"
bar.x_axis.delete = False; bar.y_axis.delete = False
db.add_chart(bar, f"B{TT+3}")
pie = PieChart(); pie.title = "Spending by category"; pie.height = 8; pie.width = 12
pie.add_data(Reference(db, min_col=12, min_row=12, max_row=12 + len(exp_cats)), titles_from_data=True)
pie.set_categories(Reference(db, min_col=11, min_row=13, max_row=12 + len(exp_cats)))
db.add_chart(pie, f"K{TT+3}")

# ---------------------------------------------------------------- Tax Estimate
tx = wb.create_sheet("Tax Estimate", 2)
setup(tx, "Tax Estimate", "Simplified US self-employment estimate for planning. Not tax advice. See Start Here.", {"A": 3, "B": 40, "C": 16, "D": 16, "E": 16, "F": 16, "G": 40})
tx["B4"] = "YEAR TO DATE"; tx["B4"].font = font(size=12, bold=True, color=TEAL)
rows = [
    (5, "Gross income", f"=Dashboard!C{TT}", USD),
    (6, "Deductible expenses", f"=SUMIFS({rng('Expenses','G')},{rng('Expenses','A')},\">=\"&DATE({y},1,1),{rng('Expenses','A')},\"<\"&DATE({y}+1,1,1))", USD),
    (7, "Mileage deduction", f"=SUMIFS({rng('Mileage','F')},{rng('Mileage','A')},\">=\"&DATE({y},1,1),{rng('Mileage','A')},\"<\"&DATE({y}+1,1,1))", USD),
    (8, "Net profit (Schedule C estimate)", "=MAX(0,C5-C6-C7)", USD),
    (10, "Profit subject to SE tax", f"=C8*{S['seshare']}", USD),
    (11, "Self-employment tax", f"=C10*{S['se']}", USD),
    (12, "Deduction for half of SE tax", "=C11/2", USD),
    (13, "Profit subject to income tax", "=MAX(0,C8-C12)", USD),
    (14, "Federal income tax", f"=C13*{S['fed']}", USD),
    (15, "State income tax", f"=C13*{S['state']}", USD),
    (17, "TOTAL ESTIMATED TAX", "=C11+C14+C15", USD),
    (18, "Effective rate on net profit", "=IF(C8>0,C17/C8,0)", PCT),
    (19, "Set aside this % of every payment", "=IF(C5>0,C17/C5,0)", PCT),
]
for r, lab, f, fmt in rows:
    tx[f"B{r}"] = lab; tx[f"C{r}"] = f; tx[f"C{r}"].number_format = fmt
    big = r in (17, 19)
    tx[f"B{r}"].font = BOLD if big or r == 8 else BODY
    tx[f"C{r}"].font = font(size=12 if big else 10, bold=big or r == 8, color=NAVY if big else "000000")
    tx[f"B{r}"].border = tx[f"C{r}"].border = BOX
    if big: tx[f"B{r}"].fill = tx[f"C{r}"].fill = LFILL
tx["C11"].comment = Comment("Social Security part stops above the yearly wage base (about $176,100 in 2025). Not modeled here.", "Hustle Ledger")

tx["B22"] = "QUARTERLY ESTIMATED PAYMENTS (IRS Form 1040-ES)"; tx["B22"].font = font(size=12, bold=True, color=TEAL)
header_row = 23
for i, lab in enumerate(["Period", "Due date", "Profit in period", "Suggested payment", "You paid", "Still owed"]):
    c = tx.cell(row=header_row, column=2 + i, value=lab); c.font, c.fill, c.border = HEAD, HFILL, BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
tx.row_dimensions[header_row].height = 30
periods = [("Q1: Jan 1 - Mar 31", f"=DATE({y},4,15)", (1, 3)), ("Q2: Apr 1 - May 31", f"=DATE({y},6,15)", (4, 5)),
           ("Q3: Jun 1 - Aug 31", f"=DATE({y},9,15)", (6, 8)), ("Q4: Sep 1 - Dec 31", f"=DATE({y}+1,1,15)", (9, 12))]
for i, (lab, due, (m0, m1)) in enumerate(periods):
    r = header_row + 1 + i
    tx[f"B{r}"] = lab; tx[f"C{r}"] = due; tx[f"C{r}"].number_format = DATEF
    tx[f"D{r}"] = f"=SUM(Dashboard!F{T0+m0}:F{T0+m1})"
    tx[f"E{r}"] = f"=MAX(0,D{r})*$C$18"
    tx[f"F{r}"].fill = IFILL; tx[f"F{r}"].font = font(size=10, color="0000FF")
    tx[f"G{r}"] = f"=MAX(0,E{r}-N(F{r}))"
    for col in "BCDEFG":
        tx[f"{col}{r}"].border = tx[f"{col}{r}"].border = BOX
        if col != "F": tx[f"{col}{r}"].font = BODY
        if col in "DEFG": tx[f"{col}{r}"].number_format = USD
tx.column_dimensions["G"].width = 16
tx["B29"] = "Due dates that fall on a weekend or holiday move to the next business day. Enter what you actually paid in the yellow cells."
tx["B29"].font = SUB

# ---------------------------------------------------------------- example / demo data
def put_income(r, d, src, desc, cat, amt, via):
    for col, v in zip("ABCDEF", (d, src, desc, cat, amt, via)): inc[f"{col}{r}"] = v
def put_exp(r, d, ven, desc, cat, amt, use=None):
    for col, v in zip("ABCDE", (d, ven, desc, cat, amt)): exp[f"{col}{r}"] = v
    if use is not None: exp[f"F{r}"] = use
def put_mile(r, d, a, b, why, miles):
    for col, v in zip("ABCDE", (d, a, b, why, miles)): mi[f"{col}{r}"] = v

if not DEMO:
    put_income(R0, date(2026, 1, 15), "Acme Co.", "Logo design (example - delete me)", "Client work", 450, "PayPal")
    put_exp(R0, date(2026, 1, 9), "Adobe", "Creative Cloud (example - delete me)", "Software & subscriptions", 59.99)
    put_mile(R0, date(2026, 1, 20), "Home", "Client office", "Kickoff meeting (example - delete me)", 18.4)
    for ws in (inc, exp, mi):
        for c in ws[R0]: c.font = font(size=10, italic=True, color=GREY)
else:
    random.seed(7)
    clients = ["Acme Co.", "Brightside Bakery", "Nova Fitness", "Etsy payout", "Fiverr payout", "Pine & Oak Realty", "Local Market Co."]
    ri = re = rm = R0
    for m in range(1, 10):
        for _ in range(random.randint(4, 7)):
            d = date(2026, m, 1 + _ * 4 + random.randint(0, 3))
            cl = random.choice(clients)
            put_income(ri, d, cl, "Project payment", "Product sales" if "payout" in cl else "Client work",
                       round(random.uniform(120, 800) * (1 + m * 0.06), 2), random.choice(pays)); ri += 1
        for cat, lo, hi in [("Software & subscriptions", 30, 90), ("Advertising", 60, 300), ("Supplies & materials", 20, 180),
                            ("Platform & payment fees", 40, 120), ("Phone & internet", 80, 80), ("Meals (business)", 25, 90)]:
            if random.random() < 0.8:
                put_exp(re, date(2026, m, random.randint(1, 28)), "Vendor", "Business purchase", cat,
                        round(random.uniform(lo, hi), 2), 0.6 if cat == "Phone & internet" else None); re += 1
        for _ in range(random.randint(1, 3)):
            put_mile(rm, date(2026, m, random.randint(1, 28)), "Home", "Client", "Client meeting", round(random.uniform(6, 40), 1)); rm += 1

wb.move_sheet("Lists", offset=10)
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
wb.active = 0
wb.save(OUT)
print("saved", OUT)
